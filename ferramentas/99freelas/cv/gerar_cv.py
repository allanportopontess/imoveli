"""Gera o CV Europass do Allan em PT-PT e EN (DOCX editável + PDF).
Contatos vêm de dados_pessoais.json (fora do git). Uso: python3 gerar_cv.py"""
import json, os, html, subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__))
P = json.load(open(os.path.join(AQUI, 'dados_pessoais.json')))
CHROME = '/opt/pw-browsers/chromium-1194/chrome-linux/chrome'
AZUL_HEX, CINZA_HEX = '0B3C8C', '5F6670'
AZUL, CINZA = RGBColor.from_string(AZUL_HEX), RGBColor.from_string(CINZA_HEX)

# ============================================
# CONTEÚDO (só fatos confirmados pelo Allan / perfil.json)
T = {
'pt': dict(
 arq='CV_Europass_Allan_Porto_PT', lang='pt-PT',
 cargo='Arquiteto e Urbanista · Arquitetura Comercial e Retail Design',
 registo='Registo profissional CAU-PE A166156-6 (Conselho de Arquitetura e Urbanismo do Brasil)',
 contato=[('Morada', 'Caruaru, Pernambuco, Brasil — disponível para relocalização em Portugal'),
          ('Telemóvel', P['telefone']), ('E-mail', P['email']),
          ('Portefólio', 'Instagram @allanportopontes_arquitetura'),
          ('Nacionalidade', 'Brasileira'), ('Carta de condução', 'Categorias A e B (carta brasileira AB)')],
 sec=dict(perfil='SOBRE MIM', exp='EXPERIÊNCIA PROFISSIONAL', edu='EDUCAÇÃO E FORMAÇÃO', cur='FORMAÇÃO COMPLEMENTAR',
          lang='COMPETÊNCIAS LINGUÍSTICAS', dig='COMPETÊNCIAS DIGITAIS', out='INFORMAÇÕES ADICIONAIS'),
 perfil='Arquiteto e urbanista com experiência nos setores público e privado — construtoras, ateliês de arquitetura e agência de publicidade. '
        'Especialista em arquitetura comercial e retail design, integro marketing, UX e UI Design na conceção de espaços estratégicos, '
        'funcionais e centrados nas pessoas. Percurso que combina projeto, gestão de projetos, gestão de obras comerciais e expansão de '
        'negócios, com forte orientação para o cliente, para prazos e para o resultado do espaço construído.',
 exp=[('2022 – atual', 'Diretor de Expansão e Novos Negócios', 'Porto Office Coworking', 'Pernambuco, Brasil',
       ['Expansão e adaptação de espaços de trabalho partilhado, com foco na experiência do utilizador e na rentabilidade.',
        'Conceção de layouts e acompanhamento da abertura e adaptação de unidades.',
        'Prospeção e desenvolvimento de novos negócios e parcerias.']),
      ('2017 – 2022', 'Arquiteto e Urbanista Sénior', 'Layout Arquitetura e Urbanismo', 'Pernambuco, Brasil',
       ['Projetos de arquitetura e de interiores residenciais e comerciais, do programa de necessidades ao acompanhamento de obra.',
        'Projetos de espaços comerciais e lojas orientados para a experiência de compra e a identidade da marca.',
        'Estudos prévios, modelação 3D e apresentações a clientes (AutoCAD, SketchUp, Layout).']),
      ('2017', 'Gestor de Obras Comerciais', 'Athena Engenharia', 'Pernambuco, Brasil',
       ['Gestão e acompanhamento da execução de obras comerciais: prazos, qualidade e conformidade com o projeto.',
        'Coordenação de empreiteiros e fornecedores em obra.']),
      ('2010 – 2014', 'Gerente de Projetos Arquitetónicos', 'URB Caruaru', 'Caruaru, Pernambuco, Brasil',
       ['Coordenação de projetos arquitetónicos e urbanos no setor público municipal.',
        'Gestão de prazos, equipas e documentação técnica; articulação com entidades públicas.'])],
 edu=[('', 'Pós-graduação em Arquitetura Comercial — Master em Retail Design', 'IPOG', 'Brasil'),
      ('2016', 'Bacharelado em Arquitetura e Urbanismo', 'UNIFAVIP DeVry', 'Caruaru, Brasil')],
 cursos=[('2025', 'Design Centrado no Utilizador', 'Escola Politécnica — PUCRS'),
         ('2025', 'Processos de UX e UI Design', 'Escola Politécnica — PUCRS'),
         ('2022', 'Foresight — Cidades Inteligentes', 'CESAR School'),
         ('2016', 'Ferramentas de Georreferenciamento — Google Earth', ''),
         ('2012', 'Análise e Gestão de Cidades', 'SEURB'),
         ('', 'Gestão de Projetos Arquitetónicos', '')],
 lang_mae=('Língua materna', 'Português'),
 lang_hdr=['', 'Compreensão oral', 'Leitura', 'Interação oral', 'Produção oral', 'Escrita'],
 lang_row=['Inglês', 'B2', 'B2', 'B2', 'B2', 'B2'],
 lang_nota='Níveis do Quadro Europeu Comum de Referência para as Línguas (QECR): A1/A2 utilizador elementar · B1/B2 utilizador independente · C1/C2 utilizador proficiente.',
 dig=[('Desenho e modelação', 'AutoCAD (avançado) · SketchUp (avançado) · Layout · Revit (básico)'),
      ('Imagem e apresentação', 'Adobe Photoshop · IA generativa (Google Gemini) para renderizações e apresentações'),
      ('Produtividade', 'Microsoft Office — Word, Excel e PowerPoint'),
      ('Tecnologia', 'Claude Code e GitHub no desenvolvimento de plataformas digitais')],
 out=['Disponibilidade imediata para relocalização para Portugal com a família.',
      'Reconhecimento do grau académico e inscrição na Ordem dos Arquitectos ao abrigo do acordo CAU–OA a iniciar.',
      'Referências profissionais disponíveis mediante pedido.']),
'en': dict(
 arq='CV_Europass_Allan_Porto_EN', lang='en',
 cargo='Architect and Urban Planner · Commercial Architecture & Retail Design',
 registo='Professional registration CAU-PE A166156-6 (Brazilian Council of Architecture and Urbanism)',
 contato=[('Address', 'Caruaru, Pernambuco, Brazil — available to relocate to Portugal'),
          ('Mobile', P['telefone']), ('E-mail', P['email']),
          ('Portfolio', 'Instagram @allanportopontes_arquitetura'),
          ('Nationality', 'Brazilian'), ('Driving licence', 'Categories A and B (Brazilian licence AB)')],
 sec=dict(perfil='ABOUT ME', exp='WORK EXPERIENCE', edu='EDUCATION AND TRAINING', cur='ADDITIONAL TRAINING',
          lang='LANGUAGE SKILLS', dig='DIGITAL SKILLS', out='ADDITIONAL INFORMATION'),
 perfil='Architect and urban planner with experience in the public and private sectors — construction companies, architecture studios and an '
        'advertising agency. Specialised in commercial architecture and retail design, I combine marketing, UX and UI Design to create '
        'strategic, functional and people-centred spaces. My background brings together design, project management, commercial '
        'construction management and business expansion, with a strong focus on clients, deadlines and the performance of the built space.',
 exp=[('2022 – present', 'Director of Expansion and New Business', 'Porto Office Coworking', 'Pernambuco, Brazil',
       ['Expansion and adaptation of shared workspaces, focused on user experience and profitability.',
        'Layout design and follow-up of the opening and refurbishment of units.',
        'Prospecting and developing new business and partnerships.']),
      ('2017 – 2022', 'Senior Architect and Urban Planner', 'Layout Arquitetura e Urbanismo', 'Pernambuco, Brazil',
       ['Residential and commercial architecture and interior design projects, from client brief to construction follow-up.',
        'Retail and store projects focused on the shopping experience and brand identity.',
        'Concept studies, 3D modelling and client presentations (AutoCAD, SketchUp, Layout).']),
      ('2017', 'Commercial Construction Manager', 'Athena Engenharia', 'Pernambuco, Brazil',
       ['Management and supervision of commercial construction works: schedules, quality and compliance with the design.',
        'Coordination of contractors and suppliers on site.']),
      ('2010 – 2014', 'Architectural Projects Manager', 'URB Caruaru', 'Caruaru, Pernambuco, Brazil',
       ['Coordination of architectural and urban projects in the municipal public sector.',
        'Management of schedules, teams and technical documentation; liaison with public bodies.'])],
 edu=[('', 'Postgraduate Degree in Commercial Architecture — Master in Retail Design', 'IPOG', 'Brazil'),
      ('2016', 'Bachelor’s Degree in Architecture and Urban Planning', 'UNIFAVIP DeVry', 'Caruaru, Brazil')],
 cursos=[('2025', 'User-Centred Design', 'Polytechnic School — PUCRS'),
         ('2025', 'UX and UI Design Processes', 'Polytechnic School — PUCRS'),
         ('2022', 'Foresight — Smart Cities', 'CESAR School'),
         ('2016', 'Georeferencing Tools — Google Earth', ''),
         ('2012', 'City Analysis and Management', 'SEURB'),
         ('', 'Architectural Project Management', '')],
 lang_mae=('Mother tongue', 'Portuguese'),
 lang_hdr=['', 'Listening', 'Reading', 'Spoken interaction', 'Spoken production', 'Writing'],
 lang_row=['English', 'B2', 'B2', 'B2', 'B2', 'B2'],
 lang_nota='Levels of the Common European Framework of Reference for Languages (CEFR): A1/A2 basic user · B1/B2 independent user · C1/C2 proficient user.',
 dig=[('Drawing and modelling', 'AutoCAD (advanced) · SketchUp (advanced) · Layout · Revit (basic)'),
      ('Imaging and presentation', 'Adobe Photoshop · Generative AI (Google Gemini) for renderings and presentations'),
      ('Productivity', 'Microsoft Office — Word, Excel and PowerPoint'),
      ('Technology', 'Claude Code and GitHub for building digital platforms')],
 out=['Immediately available to relocate to Portugal with my family.',
      'Degree recognition and registration with the Ordem dos Arquitectos under the CAU–OA agreement to be started.',
      'Professional references available upon request.']),
}

# ============================================
# DOCX
def sem_bordas(t):
  b = OxmlElement('w:tblBorders')
  for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
    x = OxmlElement('w:' + e); x.set(qn('w:val'), 'nil'); b.append(x)
  t._tbl.tblPr.append(b)

def borda_cel(cel, cor='BBBBBB'):
  tcPr = cel._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
  for e in ('top', 'left', 'bottom', 'right'):
    x = OxmlElement('w:' + e)
    for k, v in (('val', 'single'), ('sz', '4'), ('color', cor)): x.set(qn('w:' + k), v)
    b.append(x)
  tcPr.append(b)

def run(p, txt, size=10, bold=False, cor=None, ital=False):
  r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold; r.italic = ital
  r.font.name = 'Arial'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
  if cor: r.font.color.rgb = cor
  return r

def para(alvo, antes=0, depois=2):
  p = alvo.add_paragraph(); f = p.paragraph_format
  f.space_before = Pt(antes); f.space_after = Pt(depois); f.line_spacing = 1.1
  return p

def titulo_sec(doc, txt):
  p = para(doc, 12, 5); run(p, txt, 10.5, True, AZUL)
  pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
  for k, v in (('val', 'single'), ('sz', '8'), ('space', '2'), ('color', AZUL_HEX)): bt.set(qn('w:' + k), v)
  bd.append(bt); pPr.append(bd)

def tabela2(doc):
  t = doc.add_table(rows=0, cols=2); sem_bordas(t); t.alignment = WD_TABLE_ALIGNMENT.LEFT
  return t

def linha2(t, esq):
  c = t.add_row().cells; c[0].width = Cm(3.6); c[1].width = Cm(13.8)
  c[0].paragraphs[0].paragraph_format.space_after = Pt(2)
  if esq: run(c[0].paragraphs[0], esq, 9, False, CINZA)
  return c[1]

def primeira(cel):
  p = cel.paragraphs[0]; f = p.paragraph_format; f.space_after = Pt(1); f.line_spacing = 1.1
  return p

def gera_docx(d, saida):
  doc = Document()
  for s in doc.sections:
    s.page_height, s.page_width = Cm(29.7), Cm(21)
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.4)
  st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
  p = para(doc, 0, 0); run(p, 'europass', 13, True, AZUL)
  p = para(doc, 6, 0); run(p, 'Allan Porto', 22, True)
  p = para(doc, 0, 1); run(p, d['cargo'], 11, False, AZUL)
  p = para(doc, 0, 6); run(p, d['registo'], 9, False, CINZA)
  t = tabela2(doc)
  for k, v in d['contato']:
    run(primeira(linha2(t, k)), v, 9.5)
  titulo_sec(doc, d['sec']['perfil'])
  p = para(doc, 0, 2); p.paragraph_format.alignment = 3; run(p, d['perfil'], 10)
  titulo_sec(doc, d['sec']['exp'])
  t = tabela2(doc)
  for dt, cg, emp, loc, itens in d['exp']:
    c = linha2(t, dt)
    run(primeira(c), cg, 10.5, True)
    p = para(c, 0, 2); run(p, emp, 9.5, True, AZUL); run(p, '  ·  ' + loc, 9, False, CINZA)
    for x in itens:
      p = para(c, 0, 1); p.paragraph_format.left_indent = Cm(0.35); p.paragraph_format.first_line_indent = Cm(-0.35)
      run(p, '•  ' + x, 9.5)
    para(c, 0, 4)
  titulo_sec(doc, d['sec']['edu'])
  t = tabela2(doc)
  for dt, tit, inst, loc in d['edu']:
    c = linha2(t, dt)
    run(primeira(c), tit, 10.5, True)
    p = para(c, 0, 5); run(p, inst, 9.5, True, AZUL); run(p, '  ·  ' + loc, 9, False, CINZA)
  titulo_sec(doc, d['sec']['cur'])
  t = tabela2(doc)
  for dt, tit, inst in d['cursos']:
    p = primeira(linha2(t, dt)); run(p, tit, 9.5, True)
    if inst: run(p, '  ·  ' + inst, 9, False, CINZA)
  titulo_sec(doc, d['sec']['lang'])
  p = para(doc, 0, 4); run(p, d['lang_mae'][0] + ': ', 9.5, False, CINZA); run(p, d['lang_mae'][1], 10, True)
  t = doc.add_table(rows=2, cols=6)
  for i, (h, v) in enumerate(zip(d['lang_hdr'], d['lang_row'])):
    for lin, txt, kw in ((0, h, dict(size=8.5, bold=True, cor=AZUL)), (1, v, dict(size=10, bold=True))):
      cel = t.rows[lin].cells[i]; borda_cel(cel)
      pp = primeira(cel); pp.alignment = 1; pp.paragraph_format.space_before = Pt(2); run(pp, txt, **kw)
  p = para(doc, 3, 0); run(p, d['lang_nota'], 8, False, CINZA, True)
  titulo_sec(doc, d['sec']['dig'])
  t = tabela2(doc)
  for k, v in d['dig']:
    run(primeira(linha2(t, k)), v, 9.5)
  titulo_sec(doc, d['sec']['out'])
  for x in d['out']:
    p = para(doc, 0, 2); p.paragraph_format.left_indent = Cm(0.35); p.paragraph_format.first_line_indent = Cm(-0.35)
    run(p, '•  ' + x, 9.5)
  doc.save(saida)

# ============================================
# HTML / PDF
CSS = """@page{size:A4;margin:13mm 15mm 13mm 15mm}
*{box-sizing:border-box}body{font-family:Arial,'Liberation Sans',Helvetica,sans-serif;color:#1d2228;margin:0;font-size:9.6pt;line-height:1.38}
.top{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:3px solid #0B3C8C;padding-bottom:9px;margin-bottom:10px}
h1{font-size:23pt;margin:0;letter-spacing:.2px}.cargo{color:#0B3C8C;font-size:11.2pt;margin-top:2px}.reg{color:#5F6670;font-size:8.6pt;margin-top:3px}
.logo{color:#0B3C8C;font-weight:bold;font-size:14pt;letter-spacing:.3px;text-align:right}.logo small{display:block;font-weight:normal;color:#5F6670;font-size:7.8pt;letter-spacing:0}
table{border-collapse:collapse;width:100%}td{vertical-align:top;padding:1.6px 0}
td.k{width:33mm;color:#5F6670;font-size:8.6pt;padding-right:6px}
h2{color:#0B3C8C;font-size:10.2pt;letter-spacing:.6px;border-bottom:1.3px solid #0B3C8C;padding-bottom:2px;margin:13px 0 6px;break-after:avoid}
.perfil{text-align:justify;margin:0}
tr.item{break-inside:avoid}tr.item td{padding-bottom:7px}
.cg{font-weight:bold;font-size:10.4pt}.emp{color:#0B3C8C;font-weight:bold;font-size:9.4pt}.loc{color:#5F6670;font-weight:normal;font-size:8.8pt}
ul{margin:3px 0 0 0;padding-left:13px}li{margin-bottom:1.5px}li::marker{color:#0B3C8C}
tr.curso td{padding:1px 0}.ct{font-weight:bold}.ci{color:#5F6670;font-size:8.8pt}
.mae{margin-bottom:5px}.mae span{color:#5F6670;font-size:8.6pt}
table.lg{width:auto}table.lg th,table.lg td{border:1px solid #c9ced6;padding:3px 10px;text-align:center}
table.lg th{color:#0B3C8C;font-size:8.2pt;background:#eef2f9}table.lg td{font-weight:bold;font-size:9.6pt}
.nota{color:#5F6670;font-size:7.8pt;font-style:italic;margin-top:4px}
.blk{break-inside:avoid}"""

def gera_html(d):
  e = html.escape
  h = [f'<!doctype html><html lang="{d["lang"]}"><head><meta charset="utf-8"><title>Allan Porto — CV Europass</title><style>{CSS}</style></head><body>',
       f'<div class="top"><div><h1>Allan Porto</h1><div class="cargo">{e(d["cargo"])}</div><div class="reg">{e(d["registo"])}</div></div>',
       '<div class="logo">europass<small>Curriculum Vitae</small></div></div><table>']
  for k, v in d['contato']:
    vv = f'<a href="mailto:{e(v)}" style="color:inherit;text-decoration:none">{e(v)}</a>' if '@' in v and ' ' not in v else e(v)
    h.append(f'<tr><td class="k">{e(k)}</td><td>{vv}</td></tr>')
  h.append(f'</table><h2>{e(d["sec"]["perfil"])}</h2><p class="perfil">{e(d["perfil"])}</p><h2>{e(d["sec"]["exp"])}</h2><table>')
  for dt, cg, emp, loc, itens in d['exp']:
    h.append(f'<tr class="item"><td class="k">{e(dt)}</td><td><div class="cg">{e(cg)}</div><div class="emp">{e(emp)} <span class="loc">· {e(loc)}</span></div><ul>'
             + ''.join(f'<li>{e(x)}</li>' for x in itens) + '</ul></td></tr>')
  h.append(f'</table><h2>{e(d["sec"]["edu"])}</h2><table>')
  for dt, tit, inst, loc in d['edu']:
    h.append(f'<tr class="item"><td class="k">{e(dt)}</td><td><div class="cg">{e(tit)}</div><div class="emp">{e(inst)} <span class="loc">· {e(loc)}</span></div></td></tr>')
  h.append(f'</table><div class="blk"><h2>{e(d["sec"]["cur"])}</h2><table>')
  for dt, tit, inst in d['cursos']:
    h.append(f'<tr class="curso"><td class="k">{e(dt)}</td><td><span class="ct">{e(tit)}</span>' + (f' <span class="ci">· {e(inst)}</span>' if inst else '') + '</td></tr>')
  h.append(f'</table></div><div class="blk"><h2>{e(d["sec"]["lang"])}</h2><div class="mae"><span>{e(d["lang_mae"][0])}:</span> <b>{e(d["lang_mae"][1])}</b></div>'
           '<table class="lg"><tr>' + ''.join(f'<th>{e(x)}</th>' for x in d['lang_hdr']) + '</tr><tr>'
           + ''.join(f'<td>{e(x)}</td>' for x in d['lang_row']) + f'</tr></table><div class="nota">{e(d["lang_nota"])}</div></div>')
  h.append(f'<div class="blk"><h2>{e(d["sec"]["dig"])}</h2><table>' + ''.join(f'<tr><td class="k">{e(k)}</td><td>{e(v)}</td></tr>' for k, v in d['dig']) + '</table></div>')
  h.append(f'<div class="blk"><h2>{e(d["sec"]["out"])}</h2><ul>' + ''.join(f'<li>{e(x)}</li>' for x in d['out']) + '</ul></div></body></html>')
  return ''.join(h)

for lg, d in T.items():
  base = os.path.join(AQUI, d['arq'])
  gera_docx(d, base + '.docx')
  open(base + '.html', 'w').write(gera_html(d))
  subprocess.run([CHROME, '--headless', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer',
                  f'--print-to-pdf={base}.pdf', f'file://{base}.html'], capture_output=True)
  print(lg, os.path.exists(base + '.pdf'))
